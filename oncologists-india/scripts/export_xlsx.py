"""Excel для клиента: онкологи Индии + лист с пояснениями."""
import csv
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
rows = list(csv.DictReader(open(ROOT / "out" / "india_oncologists.csv", encoding="utf-8")))
(ROOT / "results").mkdir(exist_ok=True)

wb = Workbook()
ws = wb.active
ws.title = "Oncologists India"
cols = list(rows[0].keys())
ws.append(cols)
for r in rows:
    ws.append([int(r[c]) if c in ("n_pubs", "last_pub_year", "email_year") and r[c].isdigit() else r[c] for c in cols])
for i, c in enumerate(cols, 1):
    ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
    ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")
    ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(c) + 2))
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions

t = Counter(r["email_type"] or "(нет)" for r in rows)
active = [r for r in rows if r["last_pub_year"].isdigit() and int(r["last_pub_year"]) >= 2020 and not r["possible_non_physician"]]
info = wb.create_sheet("Пояснения")
for a, b in [
    ("Онкологи Индии (медицинская онкология, онкогематология, детская онкология)", ""),
    ("Всего людей", len(rows)),
    ("С имейлом", sum(1 for r in rows if r["email"])),
    ("Публиковались с 2020 года (скорее всего практикуют)", len(active)),
    ("  из них с имейлом", sum(1 for r in active if r["email"])),
    ("Члены ISMPO (общество медицинских и детских онкологов)", sum(1 for r in rows if r["ismpo_no"])),
    ("", ""),
    ("email_type", "Что значит"),
    ("published_personal", f"Врач сам указал этот адрес в своей статье (в адресе есть его имя/фамилия). Самые надёжные. ({t['published_personal']})"),
    ("published_unverified_owner", f"Адрес из статьи этого врача, но имени в адресе нет — может принадлежать соавтору. ({t['published_unverified_owner']})"),
    ("trial_contact", f"Контакт врача в клиническом исследовании (ClinicalTrials.gov). ({t['trial_contact']})"),
    ("trial_contact_generic", f"Общий ящик центра из ClinicalTrials.gov. ({t['trial_contact_generic']})"),
    ("", ""),
    ("ВАЖНО", "Перед рассылкой прогнать адреса через сервис проверки (MillionVerifier / ZeroBounce / NeverBounce)."),
    ("", "Адреса из старых статей (email_year) могли устареть; 74% адресов — личная почта (gmail и т.п.), она меняется редко."),
    ("", ""),
    ("Колонка", "Смысл"),
    ("specialty", "Отделение из аффилиации статей (медицинская онкология / онкогематология / детская / клин. гематология)"),
    ("institution, city, state", "Место работы по самой свежей статье"),
    ("last_pub_year, n_pubs", "Год последней статьи и число статей 2014–2026 (активность, авторитет)"),
    ("ismpo_no", "Номер члена ISMPO (Indian Society of Medical & Paediatric Oncology), публичный список на ismpo.org"),
    ("possible_non_physician", "yes = по аффилиациям похоже на медсестру/учёного/статистика, а не врача"),
    ("specialty_unclear", "yes = аффилиация «Departments of Pathology and Medical Oncology» и т.п., специальность неочевидна"),
    ("", ""),
    ("Источники", "Europe PMC (статьи и полные тексты открытого доступа), ISMPO members list, ClinicalTrials.gov. "
                  "Реестр NMC, CTRI (капча), WHO ICTRP (запрет в robots.txt), Practo/JustDial не использовались."),
]:
    info.append([a, b])
info.column_dimensions["A"].width = 55
info.column_dimensions["B"].width = 110
for row in (1, 8, 17):
    info.cell(row, 1).font = Font(bold=True)

out = ROOT / "results" / "oncologists_india.xlsx"
wb.save(out)
print(out)
