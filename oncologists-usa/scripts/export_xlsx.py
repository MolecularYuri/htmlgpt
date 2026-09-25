"""Excel для клиента: лист с врачами + лист с пояснениями по колонкам и статусам имейлов."""
import csv
from collections import Counter
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
rows = list(csv.DictReader(open(ROOT / "out" / "final_oncologists.csv", encoding="utf-8")))
NUM = {"partd_onc_claims", "pharma_companies_n", "pharma_trials_pi_n"}

wb = Workbook()
ws = wb.active
ws.title = "Oncologists"
cols = list(rows[0].keys())
ws.append(cols)
for r in rows:
    ws.append([float(r[c]) if (c in NUM or c.startswith("pharma_payments_")) and r[c] else r[c] for c in cols])
for i, c in enumerate(cols, 1):
    ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
    ws.cell(1, i).fill = PatternFill("solid", fgColor="1F4E78")
    ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(c) + 2))
ws.freeze_panes = "E2"
ws.auto_filter.ref = ws.dimensions

st = Counter(r["email_status"] or "(нет)" for r in rows)
info = wb.create_sheet("Пояснения")
lines = [
    ("Онкологи США (лекарственная терапия)", ""),
    ("Всего врачей", len(rows)),
    ("С имейлом", sum(1 for r in rows if r["email"])),
    ("", ""),
    ("email_status", "Что значит / ожидаемая точность"),
    ("published_personal", f"Врач сам указал адрес в статье PubMed или в ClinicalTrials.gov. Самые надёжные. ({st['published_personal']})"),
    ("pattern_high", f"Построен по схеме домена, где ≥90% адресов по одному шаблону. На проверке ~84% верных. ({st['pattern_high']})"),
    ("pattern_medium", f"Схема домена 70–90%. ~66% верных, запасные варианты в other_emails. ({st['pattern_medium']})"),
    ("pattern_low", f"Схема домена 50–70%. ~48% верных, запасные варианты в other_emails. ({st['pattern_low']})"),
    ("published_unverified_owner", f"Адрес из статьи этого врача, но в адресе нет фамилии — может быть чужой. ({st['published_unverified_owner']})"),
    ("published_research_office", f"Общий ящик исследовательского офиса из ClinicalTrials.gov. ({st['published_research_office']})"),
    ("", ""),
    ("ВАЖНО", "Все pattern_* адреса перед рассылкой прогнать через сервис проверки (MillionVerifier / ZeroBounce / NeverBounce)."),
    ("", "Старые адреса из статей (см. email_note, год публикации) тоже лучше проверить."),
    ("", ""),
    ("Колонка", "Источник"),
    ("npi, имя, специальность, адрес, телефон", "NPPES — федеральный реестр врачей (CMS), сентябрь 2026"),
    ("clinic, hospitals, in_medicare", "CMS Doctors & Clinicians + Facility Affiliation"),
    ("likely_active", "yes = в Medicare или запись в реестре обновлялась с 2021 года"),
    ("partd_onc_claims, partd_top_onc_drugs", "Medicare Part D: сколько рецептов на онкопрепараты выписал врач"),
    ("pharma_*", "Open Payments: выплаты от фармкомпаний, компании, препараты, исследования где врач — PI"),
    ("drug_signal", "yes = есть хотя бы один признак работы с лекарствами (Part D, выплаты фармы или PI исследований)"),
]
for a, b in lines:
    info.append([a, b])
info.column_dimensions["A"].width = 42
info.column_dimensions["B"].width = 110
for row in (1, 5, 16):
    info.cell(row, 1).font = Font(bold=True)
    info.cell(row, 2).font = Font(bold=True)

out = ROOT / "results" / "oncologists_usa.xlsx"
wb.save(out)
print(out)
